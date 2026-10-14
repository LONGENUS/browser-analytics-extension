"""
TailorResume — Module 8: Resume DOCX Generation Service
Generates genuine, editable Microsoft Word (.docx) documents from structured resume data.

Key Architecture:
- Consumes normalized structured resume data + template config from Module 7.
- Creates real, editable text paragraphs, runs, and list items (never images).
- Maps ATS Classic, Professional, and Compact templates into DOCX styling.
- Sets standard margins, typography sizes, colors, and line spacing.
- Embeds genuine clickable hyperlinks with protocol sanitization.
- Completely omits empty sections and handles long content gracefully.
"""

import io
from typing import Dict, Any, List, Optional
from xml.sax.saxutils import escape

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.opc.constants import RELATIONSHIP_TYPE


def sanitize_url(url: Optional[str]) -> Optional[str]:
    """Sanitizes URL schemes to allow strictly http, https, mailto, tel."""
    if not url:
        return None
    cleaned = str(url).strip()
    lower = cleaned.lower()
    if (
        lower.startswith("javascript:")
        or lower.startswith("data:")
        or lower.startswith("vbscript:")
        or lower.startswith("file:")
    ):
        return None
    if lower.startswith(("http://", "https://", "mailto:", "tel:")):
        return cleaned
    return f"https://{cleaned}"


def format_date_range(start_date: Optional[str], end_date: Optional[str], is_current: bool = False) -> str:
    """Formats start and end dates into a clean range string."""
    start = (start_date or "").strip()
    if is_current:
        end = "Present"
    else:
        end = (end_date or "").strip()

    if start and end:
        return f"{start} – {end}"
    if start:
        return start
    if end:
        return end
    return ""


def add_hyperlink(paragraph, url: str, text: str, color_hex: str = "2563EB", underline: bool = True):
    """Embeds an actual clickable Word hyperlink into a paragraph."""
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = parse_xml(
        f'<w:hyperlink {nsdecls("w")} r:id="{r_id}" {nsdecls("r")}/>'
    )
    new_run = parse_xml(
        f'<w:r {nsdecls("w")}>'
        f'<w:rPr>'
        f'<w:color w:val="{color_hex}"/>'
        f'{"<w:u w:val=\"single\"/>" if underline else ""}'
        f'</w:rPr>'
        f'<w:t>{escape(text)}</w:t>'
        f'</w:r>'
    )
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_bottom_border(paragraph, color_hex: str = "CBD5E1", size_eighths: int = 8):
    """Adds a bottom border line beneath a paragraph in Word."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'<w:bottom w:val="single" w:sz="{size_eighths}" w:space="4" w:color="{color_hex}"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)


class DOCXResumeGenerator:
    """Generates styled Word documents from structured resume data."""

    def __init__(self, resume_data: Dict[str, Any], template_id: str = "ats-classic", template_version: str = "v1.0.0"):
        self.data = resume_data
        self.template_id = (template_id or "ats-classic").lower().strip()
        if self.template_id not in ("ats-classic", "professional", "compact"):
            self.template_id = "ats-classic"
        self.template_version = template_version or "v1.0.0"

        self._init_template_tokens()

    def _init_template_tokens(self):
        """Initializes typography and margin tokens matching Module 7."""
        if self.template_id == "professional":
            self.font_name = "Calibri"
            self.margin_in = 0.8
            self.name_size = Pt(24)
            self.headline_size = Pt(11.5)
            self.heading_size = Pt(11)
            self.item_size = Pt(10.5)
            self.body_size = Pt(10)
            self.bullet_size = Pt(9.5)
            self.color_primary = RGBColor(15, 23, 42)     # #0F172A
            self.color_secondary = RGBColor(51, 65, 85)   # #334155
            self.color_muted = RGBColor(100, 116, 139)    # #64748B
            self.color_accent = RGBColor(37, 99, 235)     # #2563EB
            self.border_hex = "2563EB"
            self.border_sz = 12
            self.space_before_section = Pt(12)
            self.space_after_section = Pt(4)
            self.item_spacing = Pt(5)
            self.bullet_spacing = Pt(2)
        elif self.template_id == "compact":
            self.font_name = "Calibri"
            self.margin_in = 0.55
            self.name_size = Pt(18)
            self.headline_size = Pt(10.5)
            self.heading_size = Pt(10)
            self.item_size = Pt(9.5)
            self.body_size = Pt(9)
            self.bullet_size = Pt(8.5)
            self.color_primary = RGBColor(15, 23, 42)
            self.color_secondary = RGBColor(30, 41, 59)
            self.color_muted = RGBColor(71, 85, 105)
            self.color_accent = RGBColor(15, 23, 42)
            self.border_hex = "94A3B8"
            self.border_sz = 6
            self.space_before_section = Pt(7)
            self.space_after_section = Pt(2)
            self.item_spacing = Pt(3)
            self.bullet_spacing = Pt(1)
        else:  # ats-classic
            self.font_name = "Calibri"
            self.margin_in = 0.9
            self.name_size = Pt(22)
            self.headline_size = Pt(11)
            self.heading_size = Pt(11)
            self.item_size = Pt(10.5)
            self.body_size = Pt(10)
            self.bullet_size = Pt(9.5)
            self.color_primary = RGBColor(15, 23, 42)
            self.color_secondary = RGBColor(51, 65, 85)
            self.color_muted = RGBColor(71, 85, 105)
            self.color_accent = RGBColor(15, 23, 42)
            self.border_hex = "CBD5E1"
            self.border_sz = 8
            self.space_before_section = Pt(10)
            self.space_after_section = Pt(3)
            self.item_spacing = Pt(4.5)
            self.bullet_spacing = Pt(1.5)

    def generate(self) -> bytes:
        """Assembles the DOCX document and outputs binary bytes."""
        doc = docx.Document()

        # Set page margins
        for section in doc.sections:
            section.top_margin = Inches(self.margin_in)
            section.bottom_margin = Inches(self.margin_in)
            section.left_margin = Inches(self.margin_in)
            section.right_margin = Inches(self.margin_in)

        # 1. Header (Personal Info)
        self._build_header(doc)

        # 2. Professional Summary
        self._build_summary(doc)

        # 3. Work Experience
        self._build_experience(doc)

        # 4. Skills
        self._build_skills(doc)

        # 5. Projects
        self._build_projects(doc)

        # 6. Education
        self._build_education(doc)

        # 7. Certifications
        self._build_certifications(doc)

        # 8. Links
        self._build_links(doc)

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    def _add_section_heading(self, doc: docx.Document, title: str):
        """Adds standard section heading with bottom border and spacing."""
        p = doc.add_paragraph()
        p.paragraph_format.space_before = self.space_before_section
        p.paragraph_format.space_after = self.space_after_section
        p.paragraph_format.keep_with_next = True

        run = p.add_run(title.upper())
        run.font.name = self.font_name
        run.font.size = self.heading_size
        run.font.bold = True
        run.font.color.rgb = self.color_accent

        add_bottom_border(p, color_hex=self.border_hex, size_eighths=self.border_sz)

    def _build_header(self, doc: docx.Document):
        p_info = self.data.get("personal_info") or {}
        full_name = (p_info.get("full_name") or "").strip()
        headline = (p_info.get("headline") or "").strip()

        # Name
        if full_name:
            p_name = doc.add_paragraph()
            p_name.paragraph_format.space_before = Pt(0)
            p_name.paragraph_format.space_after = Pt(2)
            run = p_name.add_run(full_name)
            run.font.name = self.font_name
            run.font.size = self.name_size
            run.font.bold = True
            run.font.color.rgb = self.color_primary

        # Headline
        if headline:
            p_head = doc.add_paragraph()
            p_head.paragraph_format.space_before = Pt(0)
            p_head.paragraph_format.space_after = Pt(3)
            run = p_head.add_run(headline)
            run.font.name = self.font_name
            run.font.size = self.headline_size
            run.font.bold = (self.template_id == "professional")
            run.font.color.rgb = self.color_accent

        # Contact Row
        contact_items = []
        if p_info.get("email"):
            contact_items.append(("email", p_info["email"], f"mailto:{p_info['email']}"))
        if p_info.get("phone"):
            contact_items.append(("text", p_info["phone"], None))
        if p_info.get("location"):
            contact_items.append(("text", p_info["location"], None))
        if p_info.get("website"):
            contact_items.append(("link", "Website", sanitize_url(p_info["website"])))
        if p_info.get("linkedin_url"):
            contact_items.append(("link", "LinkedIn", sanitize_url(p_info["linkedin_url"])))
        if p_info.get("github_url"):
            contact_items.append(("link", "GitHub", sanitize_url(p_info["github_url"])))
        if p_info.get("portfolio_url"):
            contact_items.append(("link", "Portfolio", sanitize_url(p_info["portfolio_url"])))

        if contact_items:
            p_contact = doc.add_paragraph()
            p_contact.paragraph_format.space_before = Pt(0)
            p_contact.paragraph_format.space_after = Pt(4)

            for i, (item_type, label, url) in enumerate(contact_items):
                if i > 0:
                    sep = p_contact.add_run("   •   ")
                    sep.font.name = self.font_name
                    sep.font.size = Pt(8.5)
                    sep.font.color.rgb = self.color_muted

                if item_type in ("email", "link") and url:
                    add_hyperlink(p_contact, url, label, color_hex="2563EB", underline=False)
                else:
                    run = p_contact.add_run(label)
                    run.font.name = self.font_name
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = self.color_muted

    def _build_summary(self, doc: docx.Document):
        summary_obj = self.data.get("summary") or {}
        summary_text = (summary_obj.get("summary") or "").strip()
        if not summary_text:
            return

        self._add_section_heading(doc, "Professional Summary")
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(summary_text)
        run.font.name = self.font_name
        run.font.size = self.body_size
        run.font.color.rgb = self.color_secondary

    def _build_experience(self, doc: docx.Document):
        experiences = self.data.get("experiences") or []
        valid_exp = [e for e in experiences if e.get("company") or e.get("job_title")]
        if not valid_exp:
            return

        self._add_section_heading(doc, "Work Experience")

        for exp in valid_exp:
            # Header line: Job Title — Company [Dates]
            p_header = doc.add_paragraph()
            p_header.paragraph_format.space_before = self.item_spacing
            p_header.paragraph_format.space_after = Pt(1)
            p_header.paragraph_format.keep_with_next = True

            run_title = p_header.add_run(exp.get("job_title", ""))
            run_title.font.name = self.font_name
            run_title.font.size = self.item_size
            run_title.font.bold = True
            run_title.font.color.rgb = self.color_primary

            if exp.get("company"):
                run_comp = p_header.add_run(f" — {exp['company']}")
                run_comp.font.name = self.font_name
                run_comp.font.size = self.item_size
                run_comp.font.italic = True
                run_comp.font.color.rgb = self.color_secondary

            date_str = format_date_range(exp.get("start_date"), exp.get("end_date"), exp.get("is_current", False))
            if date_str:
                run_date = p_header.add_run(f"  ({date_str})")
                run_date.font.name = self.font_name
                run_date.font.size = Pt(9)
                run_date.font.color.rgb = self.color_muted

            # Sub-meta: Location & Employment type
            meta_parts = [exp.get("location"), exp.get("employment_type")]
            meta_clean = [m for m in meta_parts if m and m.strip()]
            if meta_clean:
                p_meta = doc.add_paragraph()
                p_meta.paragraph_format.space_before = Pt(0)
                p_meta.paragraph_format.space_after = Pt(1.5)
                p_meta.paragraph_format.keep_with_next = True
                run_meta = p_meta.add_run(" • ".join(meta_clean))
                run_meta.font.name = self.font_name
                run_meta.font.size = Pt(8.5)
                run_meta.font.color.rgb = self.color_muted

            # Description
            if exp.get("description"):
                p_desc = doc.add_paragraph()
                p_desc.paragraph_format.space_before = Pt(1)
                p_desc.paragraph_format.space_after = Pt(2)
                run_desc = p_desc.add_run(exp["description"])
                run_desc.font.name = self.font_name
                run_desc.font.size = self.body_size
                run_desc.font.color.rgb = self.color_secondary

            # Bullets
            bullets = exp.get("bullets") or []
            for b in bullets:
                content = b.get("content") if isinstance(b, dict) else str(b)
                if content and content.strip():
                    p_bullet = doc.add_paragraph(style="List Bullet")
                    p_bullet.paragraph_format.space_before = Pt(0)
                    p_bullet.paragraph_format.space_after = self.bullet_spacing
                    run_b = p_bullet.add_run(content.strip())
                    run_b.font.name = self.font_name
                    run_b.font.size = self.bullet_size
                    run_b.font.color.rgb = self.color_secondary

    def _build_skills(self, doc: docx.Document):
        skills = self.data.get("skills") or []
        valid_skills = [s for s in skills if s.get("skill_name") and s.get("skill_name").strip()]
        if not valid_skills:
            return

        self._add_section_heading(doc, "Skills")

        has_categories = any(s.get("category") and s.get("category").strip() for s in valid_skills)
        if has_categories:
            categorized = {}
            for s in valid_skills:
                cat = (s.get("category") or "General").strip()
                categorized.setdefault(cat, []).append(s["skill_name"].strip())

            for cat, s_list in categorized.items():
                p_cat = doc.add_paragraph()
                p_cat.paragraph_format.space_before = Pt(1)
                p_cat.paragraph_format.space_after = Pt(1.5)

                run_cat = p_cat.add_run(f"{cat}: ")
                run_cat.font.name = self.font_name
                run_cat.font.size = self.body_size
                run_cat.font.bold = True
                run_cat.font.color.rgb = self.color_primary

                sep = " • " if self.template_id in ("professional", "compact") else ", "
                run_skills = p_cat.add_run(sep.join(s_list))
                run_skills.font.name = self.font_name
                run_skills.font.size = self.body_size
                run_skills.font.color.rgb = self.color_secondary
        else:
            p_flat = doc.add_paragraph()
            p_flat.paragraph_format.space_before = Pt(1)
            p_flat.paragraph_format.space_after = Pt(2)
            run_flat = p_flat.add_run(", ".join(s["skill_name"].strip() for s in valid_skills))
            run_flat.font.name = self.font_name
            run_flat.font.size = self.body_size
            run_flat.font.color.rgb = self.color_secondary

    def _build_projects(self, doc: docx.Document):
        projects = self.data.get("projects") or []
        valid_projects = [p for p in projects if p.get("name") and p.get("name").strip()]
        if not valid_projects:
            return

        self._add_section_heading(doc, "Projects")

        for proj in valid_projects:
            p_proj = doc.add_paragraph()
            p_proj.paragraph_format.space_before = self.item_spacing
            p_proj.paragraph_format.space_after = Pt(1)
            p_proj.paragraph_format.keep_with_next = True

            run_name = p_proj.add_run(proj.get("name", ""))
            run_name.font.name = self.font_name
            run_name.font.size = self.item_size
            run_name.font.bold = True
            run_name.font.color.rgb = self.color_primary

            safe_url = sanitize_url(proj.get("url"))
            if safe_url:
                p_proj.add_run("  ")
                add_hyperlink(p_proj, safe_url, "[Link]", color_hex="2563EB", underline=False)

            if proj.get("technologies"):
                run_tech = p_proj.add_run(f"  ({proj['technologies']})")
                run_tech.font.name = self.font_name
                run_tech.font.size = Pt(8.5)
                run_tech.font.color.rgb = self.color_muted

            if proj.get("description"):
                p_desc = doc.add_paragraph()
                p_desc.paragraph_format.space_before = Pt(1)
                p_desc.paragraph_format.space_after = Pt(2)
                run_desc = p_desc.add_run(proj["description"])
                run_desc.font.name = self.font_name
                run_desc.font.size = self.body_size
                run_desc.font.color.rgb = self.color_secondary

            bullets = proj.get("bullets") or []
            for b in bullets:
                content = b.get("content") if isinstance(b, dict) else str(b)
                if content and content.strip():
                    p_b = doc.add_paragraph(style="List Bullet")
                    p_b.paragraph_format.space_before = Pt(0)
                    p_b.paragraph_format.space_after = self.bullet_spacing
                    run_b = p_b.add_run(content.strip())
                    run_b.font.name = self.font_name
                    run_b.font.size = self.bullet_size
                    run_b.font.color.rgb = self.color_secondary

    def _build_education(self, doc: docx.Document):
        educations = self.data.get("educations") or []
        valid_edu = [e for e in educations if e.get("institution") and e.get("institution").strip()]
        if not valid_edu:
            return

        self._add_section_heading(doc, "Education")

        for edu in valid_edu:
            p_edu = doc.add_paragraph()
            p_edu.paragraph_format.space_before = self.item_spacing
            p_edu.paragraph_format.space_after = Pt(1)
            p_edu.paragraph_format.keep_with_next = True

            run_inst = p_edu.add_run(edu.get("institution", ""))
            run_inst.font.name = self.font_name
            run_inst.font.size = self.item_size
            run_inst.font.bold = True
            run_inst.font.color.rgb = self.color_primary

            if edu.get("location"):
                run_loc = p_edu.add_run(f" — {edu['location']}")
                run_loc.font.name = self.font_name
                run_loc.font.size = self.item_size
                run_loc.font.color.rgb = self.color_secondary

            date_str = format_date_range(edu.get("start_date"), edu.get("end_date"), False)
            if date_str:
                run_date = p_edu.add_run(f"  ({date_str})")
                run_date.font.name = self.font_name
                run_date.font.size = Pt(8.5)
                run_date.font.color.rgb = self.color_muted

            deg_parts = [edu.get("degree"), edu.get("field_of_study")]
            deg_clean = [d for d in deg_parts if d and d.strip()]
            if deg_clean:
                p_deg = doc.add_paragraph()
                p_deg.paragraph_format.space_before = Pt(0)
                p_deg.paragraph_format.space_after = Pt(1.5)
                run_deg = p_deg.add_run(" in ".join(deg_clean))
                run_deg.font.name = self.font_name
                run_deg.font.size = self.body_size
                run_deg.font.color.rgb = self.color_secondary

            if edu.get("description"):
                p_desc = doc.add_paragraph()
                p_desc.paragraph_format.space_before = Pt(1)
                p_desc.paragraph_format.space_after = Pt(2)
                run_desc = p_desc.add_run(edu["description"])
                run_desc.font.name = self.font_name
                run_desc.font.size = self.body_size
                run_desc.font.color.rgb = self.color_secondary

    def _build_certifications(self, doc: docx.Document):
        certs = self.data.get("certifications") or []
        valid_certs = [c for c in certs if c.get("name") and c.get("name").strip()]
        if not valid_certs:
            return

        self._add_section_heading(doc, "Certifications")

        for cert in valid_certs:
            p_cert = doc.add_paragraph()
            p_cert.paragraph_format.space_before = self.item_spacing
            p_cert.paragraph_format.space_after = Pt(1)
            p_cert.paragraph_format.keep_with_next = True

            run_name = p_cert.add_run(cert.get("name", ""))
            run_name.font.name = self.font_name
            run_name.font.size = self.item_size
            run_name.font.bold = True
            run_name.font.color.rgb = self.color_primary

            if cert.get("issuer"):
                run_iss = p_cert.add_run(f" • {cert['issuer']}")
                run_iss.font.name = self.font_name
                run_iss.font.size = self.body_size
                run_iss.font.color.rgb = self.color_secondary

            date_str = format_date_range(cert.get("issue_date"), cert.get("expiration_date"), False)
            if date_str:
                run_date = p_cert.add_run(f"  ({date_str})")
                run_date.font.name = self.font_name
                run_date.font.size = Pt(8.5)
                run_date.font.color.rgb = self.color_muted

            safe_url = sanitize_url(cert.get("credential_url"))
            if safe_url:
                p_cert.add_run("  ")
                add_hyperlink(p_cert, safe_url, "[Verify]", color_hex="2563EB", underline=False)

            if cert.get("credential_id"):
                p_id = doc.add_paragraph()
                p_id.paragraph_format.space_before = Pt(0)
                p_id.paragraph_format.space_after = Pt(1.5)
                run_id = p_id.add_run(f"Credential ID: {cert['credential_id']}")
                run_id.font.name = self.font_name
                run_id.font.size = Pt(8.5)
                run_id.font.color.rgb = self.color_muted

    def _build_links(self, doc: docx.Document):
        links = self.data.get("links") or []
        valid_links = [l for l in links if l.get("label") and sanitize_url(l.get("url"))]
        if not valid_links:
            return

        self._add_section_heading(doc, "Links & Portfolios")
        p_links = doc.add_paragraph()
        p_links.paragraph_format.space_before = Pt(1)
        p_links.paragraph_format.space_after = Pt(2)

        for i, l in enumerate(valid_links):
            if i > 0:
                sep = p_links.add_run("   •   ")
                sep.font.name = self.font_name
                sep.font.size = Pt(8.5)
                sep.font.color.rgb = self.color_muted

            safe_url = sanitize_url(l["url"])
            add_hyperlink(p_links, safe_url, l["label"], color_hex="2563EB", underline=False)


def generate_resume_docx(resume_data: Dict[str, Any], template_id: str = "ats-classic", template_version: str = "v1.0.0") -> bytes:
    """Convenience helper to generate DOCX bytes from structured resume data."""
    generator = DOCXResumeGenerator(resume_data, template_id, template_version)
    return generator.generate()
